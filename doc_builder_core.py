"""
doc_builder_core.py
===================
Módulo central de estilos, formateo avanzado, encabezados, pies de página,
cajas de texto destacadas (callouts), tablas profesionales e inserción de figuras
para el Informe Profesional de Presentación de FitExpert.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Paleta Cromática Institucional (Slate & Emerald Executive)
COLOR_NAVY_DARK = RGBColor(0x0F, 0x17, 0x2A)   # Slate 900
COLOR_NAVY_LIGHT = RGBColor(0x1E, 0x29, 0x3B)  # Slate 800
COLOR_BLUE_ACCENT = RGBColor(0x02, 0x84, 0xC7) # Sky 600
COLOR_EMERALD = RGBColor(0x05, 0x96, 0x69)     # Emerald 600
COLOR_TEXT_MAIN = RGBColor(0x1E, 0x29, 0x3B)   # Slate 800
COLOR_TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)  # Slate 500
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEX_NAVY_DARK = "0F172A"
HEX_NAVY_LIGHT = "1E293B"
HEX_EMERALD = "059669"
HEX_EMERALD_LIGHT = "ECFDF5"
HEX_BLUE_LIGHT = "EFF6FF"
HEX_BG_LIGHT = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BORDER = "0284C7"


def setup_document_styles(doc: docx.Document):
    """Configura márgenes, estilos base, tipografías y párrafos del documento."""
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.different_first_page_header_footer = True

        # Encabezado (Páginas subsiguientes)
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("FITEXPERT v2.0  |  INFORME EJECUTIVO Y DE PRESENTACIÓN")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_TEXT_MUTED

        # Pie de página (Páginas subsiguientes)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun1 = fp.add_run("Documento de Presentación Institucional y Funcional  —  ")
        frun1.font.name = "Calibri"
        frun1.font.size = Pt(9)
        frun1.font.color.rgb = COLOR_TEXT_MUTED

        # Campo de página Word XML
        fld = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(fld)

    # Normal Style
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = COLOR_TEXT_MAIN
    normal_style.paragraph_format.line_spacing = 1.18
    normal_style.paragraph_format.space_after = Pt(6)


def add_title_header(doc: docx.Document, text: str, level: int = 1):
    """Agrega un encabezado numerado con diseño tipográfico formal."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.bold = True

    if level == 1:
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(18)
        run.font.color.rgb = COLOR_NAVY_DARK
        # Añadir línea decorativa sutil
        pBdr = parse_xml(r'<w:pBdr %s><w:bottom w:val="single" w:sz="12" w:space="4" w:color="0284C7"/></w:pBdr>' % nsdecls('w'))
        p._p.get_or_add_pPr().append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(14)
        run.font.color.rgb = COLOR_NAVY_LIGHT
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_EMERALD
    elif level == 4:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_TEXT_MAIN
        run.italic = True
    return p


def add_p(doc: docx.Document, text: str, bold_prefix: str = None, space_after: int = 6):
    """Agrega un párrafo formal con justificación y espaciado consistente."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(space_after)

    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(11)
        r_pre.bold = True
        r_pre.font.color.rgb = COLOR_NAVY_DARK

    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_TEXT_MAIN
    return p


def add_bullet(doc: docx.Document, text: str, bold_prefix: str = None):
    """Agrega un elemento de lista con viñeta y tipografía uniforme."""
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)

    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.bold = True
        r_pre.font.color.rgb = COLOR_NAVY_DARK

    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_TEXT_MAIN
    return p


def add_callout(doc: docx.Document, text: str, title: str = "NOTA INFORMATIVA", border_color: str = HEX_CALLOUT_BORDER, bg_color: str = HEX_BLUE_LIGHT):
    """Crea una caja de llamada visualmente destacada (estilo alert box)."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)

    # Shading
    shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), bg_color))
    cell._tc.get_or_add_tcPr().append(shd)

    # Margins and Borders
    borders = parse_xml(r'''
        <w:tcBorders %s>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="%s"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''' % (nsdecls('w'), border_color))
    cell._tc.get_or_add_tcPr().append(borders)

    mar = parse_xml(r'''
        <w:tcMar %s>
            <w:top w:w="140" w:type="dxa"/>
            <w:left w:w="220" w:type="dxa"/>
            <w:bottom w:w="140" w:type="dxa"/>
            <w:right w:w="200" w:type="dxa"/>
        </w:tcMar>
    ''' % nsdecls('w'))
    cell._tc.get_or_add_tcPr().append(mar)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_t = p.add_run(f"📌 {title}: ")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(10.5)
    r_t.bold = True
    r_t.font.color.rgb = COLOR_NAVY_DARK

    r_txt = p.add_run(text)
    r_txt.font.name = "Calibri"
    r_txt.font.size = Pt(10)
    r_txt.font.color.rgb = COLOR_TEXT_MAIN

    # Párrafo vacío posterior para espaciado
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(6)


def add_table_data(doc: docx.Document, headers: list, data: list, col_widths: list = None, header_bg: str = HEX_NAVY_DARK):
    """Crea una tabla formateada profesionalmente con cabecera oscura y filas alternas."""
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Cabecera
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        if col_widths and i < len(col_widths):
            hdr_cells[i].width = Inches(col_widths[i])
        shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), header_bg))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)

        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.bold = True
            r.font.color.rgb = COLOR_WHITE

    # Filas de datos
    for r_idx, row_values in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_row = HEX_BG_LIGHT if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            if col_widths and c_idx < len(col_widths):
                row_cells[c_idx].width = Inches(col_widths[c_idx])
            shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), bg_row))
            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_TEXT_MAIN

    # Bordes sutiles para la tabla completa
    tblBorders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="0F172A"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tbl._tbl.tblPr.append(tblBorders)

    # Espaciador posterior
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(6)
    return tbl


def add_figure(doc: docx.Document, img_path: str, caption_text: str, width_in: float = 6.0):
    """Inserta una imagen centrada con su respectiva leyenda formal."""
    if not os.path.exists(img_path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_in))

    # Caption
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(12)
    crun = cp.add_run(caption_text)
    crun.font.name = "Calibri"
    crun.font.size = Pt(9.5)
    crun.bold = True
    crun.font.color.rgb = COLOR_TEXT_MUTED
