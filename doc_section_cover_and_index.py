"""
doc_section_cover_and_index.py
==============================
Generación de la Portada Profesional y del Índice General de Contenidos
para el Informe de Presentación del Proyecto FitExpert.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from doc_builder_core import (
    COLOR_NAVY_DARK, COLOR_NAVY_LIGHT, COLOR_BLUE_ACCENT, COLOR_EMERALD,
    COLOR_TEXT_MAIN, COLOR_TEXT_MUTED, COLOR_WHITE, HEX_NAVY_DARK,
    HEX_BG_LIGHT, add_title_header, add_p
)


def build_cover_page(doc: docx.Document):
    """Genera una portada de nivel ejecutivo/académico formal y limpia."""
    # Espaciado superior inicial
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(36)
    p_top.paragraph_format.space_after = Pt(8)

    # Super-encabezado institucional
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("INTELIGENCIA ARTIFICIAL SIMBÓLICA  ·  SISTEMAS BASADOS EN CONOCIMIENTO")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(10)
    r_inst.bold = True
    r_inst.font.color.rgb = COLOR_BLUE_ACCENT

    # Línea decorativa
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(28)
    r_line = p_line.add_run("—" * 38)
    r_line.font.color.rgb = COLOR_TEXT_MUTED

    # Título Principal del Proyecto
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("FITEXPERT")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(38)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_NAVY_DARK

    # Subtítulo Descriptivo
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run("Sistema Experto Inteligente para la Prescripción Personalizada\nde Nutrición y Acondicionamiento Físico Biomecánicamente Seguro")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(15)
    r_sub.bold = True
    r_sub.font.color.rgb = COLOR_EMERALD

    # Tipo de Documento
    p_doc = doc.add_paragraph()
    p_doc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_doc.paragraph_format.space_after = Pt(45)
    r_doc = p_doc.add_run("INFORME EJECUTIVO, CONCEPTUAL Y DE PRESENTACIÓN INTEGRAL DEL PROYECTO")
    r_doc.font.name = "Calibri"
    r_doc.font.size = Pt(11)
    r_doc.bold = True
    r_doc.font.color.rgb = COLOR_TEXT_MUTED

    # Tabla de Ficha Técnica / Metadatos en la portada
    tbl = doc.add_table(rows=8, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    metadata = [
        ("Nombre del Proyecto:", "FitExpert — Asesor Inteligente en Salud y Rendimiento"),
        ("Versión del Sistema:", "2.0 (Release Integral Académica e Industrial)"),
        ("Autor Principal:", "Oliver García (ogarcore)"),
        ("Contacto:", "garciaooliver@gmail.com"),
        ("Área de Conocimiento:", "Ciencias de la Computación / Inteligencia Artificial"),
        ("Línea de Especialización:", "Sistemas Basados en Conocimiento · Motores de Inferencia"),
        ("Entorno Tecnológico:", "Python 3 · CustomTkinter · Streamlit · Rich · ReportLab"),
        ("Fecha de Publicación:", "Octubre de 2026"),
    ]

    col_widths = [2.2, 4.3]
    for row_idx, (label, val) in enumerate(metadata):
        row_cells = tbl.rows[row_idx].cells
        for c_idx in range(2):
            row_cells[c_idx].width = Inches(col_widths[c_idx])
            shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), HEX_BG_LIGHT if row_idx % 2 == 1 else "FFFFFF"))
            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

        # Label
        p_lbl = row_cells[0].paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(3)
        p_lbl.paragraph_format.space_after = Pt(3)
        r_l = p_lbl.add_run(label)
        r_l.font.name = "Calibri"
        r_l.font.size = Pt(9.5)
        r_l.bold = True
        r_l.font.color.rgb = COLOR_NAVY_DARK

        # Value
        p_val = row_cells[1].paragraphs[0]
        p_val.paragraph_format.space_before = Pt(3)
        p_val.paragraph_format.space_after = Pt(3)
        r_v = p_val.add_run(val)
        r_v.font.name = "Calibri"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = COLOR_TEXT_MAIN

    # Bordes elegantes para la tabla de portada
    tblBorders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="0284C7"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="0284C7"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tbl._tbl.tblPr.append(tblBorders)

    # Nota inferior
    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bot.paragraph_format.space_before = Pt(45)
    r_bot = p_bot.add_run("Documentación oficial orientada a presentación conceptual, evaluación funcional y auditoría de calidad.")
    r_bot.font.name = "Calibri"
    r_bot.font.size = Pt(8.5)
    r_bot.italic = True
    r_bot.font.color.rgb = COLOR_TEXT_MUTED

    doc.add_page_break()


def build_table_of_contents(doc: docx.Document):
    """Genera la estructura del Índice General de Contenidos formal y exhaustivo."""
    add_title_header(doc, "Índice General de Contenidos", level=1)
    add_p(doc, "El presente documento ha sido estructurado con rigor profesional y académico para facilitar una comprensión integral del sistema FitExpert. A continuación se desglosan los capítulos, módulos y apartados analíticos que conforman este informe:")

    toc_items = [
        ("PORTADA INSTITUCIONAL", "1"),
        ("ÍNDICE GENERAL DE CONTENIDOS", "2"),
        ("1. INTRODUCCIÓN", "3"),
        ("   1.1 Contexto de Surgimiento y Motivación", "3"),
        ("   1.2 La Necesidad Social y Sanitaria", "4"),
        ("   1.3 Propósito del Presente Informe", "5"),
        ("2. DESCRIPCIÓN GENERAL DEL PROYECTO", "6"),
        ("   2.1 Definición Conceptual: ¿Qué es FitExpert?", "6"),
        ("   2.2 Finalidad Primordial y Entorno de Aplicación", "7"),
        ("   2.3 Actividades Principales que Facilita el Sistema", "8"),
        ("3. ANTECEDENTES Y PROBLEMA", "9"),
        ("   3.1 La Crisis del Modelo Tradicional de Asesoría", "9"),
        ("   3.2 La Trampa de los Planes Masivos y Despersonalizados", "10"),
        ("   3.3 La Epidemia Silenciosa de Lesiones Articulares", "11"),
        ("   3.4 Restricciones Alergénicas Ignoradas y Efecto Rebote", "12"),
        ("4. JUSTIFICACIÓN", "13"),
        ("   4.1 Relevancia e Impacto Social de la Solución", "13"),
        ("   4.2 Fundamento Científico y Seguridad Biomecánica", "14"),
        ("   4.3 Democratización del Acceso al Bienestar Físico", "15"),
        ("5. OBJETIVOS", "16"),
        ("   5.1 Objetivo General", "16"),
        ("   5.2 Objetivos Específicos Detallados", "16"),
        ("6. ALCANCE DEL PROYECTO", "18"),
        ("   6.1 Cobertura Funcional y Procesos Integrados", "18"),
        ("   6.2 Límites Operativos y Fronteras Médicas del Sistema", "19"),
        ("7. USUARIOS DEL SISTEMA", "20"),
        ("   7.1 Usuarios Finales y Atletas (Principiantes a Avanzados)", "20"),
        ("   7.2 Poblaciones Especiales (Seniors, Menores, Lesionados, Obesidad)", "21"),
        ("   7.3 Profesionales de la Salud y Entrenadores Personales", "22"),
        ("   7.4 Evaluadores Técnicos y Comunidad Académica", "23"),
        ("8. FUNCIONAMIENTO GENERAL", "24"),
        ("   8.1 El Ciclo Operativo de Extremo a Extremo", "24"),
        ("   8.2 Fases del Flujo de Interacción del Usuario", "25"),
        ("   8.3 Diagrama del Recorrido Funcional del Sistema", "27"),
        ("9. PRINCIPALES FUNCIONALIDADES", "28"),
        ("   9.1 Autenticación y Control de Identidad Criptográfica Local", "28"),
        ("   9.2 Diagnóstico Antropométrico y Balance Energético Automatizado", "29"),
        ("   9.3 Prescripción Nutricional Dinámica y Desglose de Macronutrientes", "31"),
        ("   9.4 Sistema de Filtrado de Alérgenos e Intolerancias Alimentarias", "32"),
        ("   9.5 Planificador Semanal de Microciclos de Entrenamiento", "34"),
        ("   9.6 Blindaje Biomecánico y Sustitución por Lesiones Articulares", "35"),
        ("   9.7 Filtros Protectores por Rango Etario y Obesidad Severa", "37"),
        ("   9.8 Módulo de Explicabilidad y Razonamiento Transparente", "38"),
        ("   9.9 Tablero de Control de Evolución Ponderal y Sobrecarga", "40"),
        ("   9.10 Generación y Exportación de Expedientes en PDF", "41"),
        ("10. RECORRIDO DEL USUARIO (CASOS REALES)", "43"),
        ("   10.1 Escenario A: Usuario Principiante en Pérdida de Grasa Doméstica", "43"),
        ("   10.2 Escenario B: Levantador Intermedio con Lesión Lumbar", "44"),
        ("   10.3 Escenario C: Adulto Mayor en Prevención de Sarcopenia", "45"),
        ("   10.4 Escenario D: Sesión de Seguimiento Mensual y Auditoría", "46"),
        ("11. PRINCIPALES MÓDULOS DEL SISTEMA", "47"),
        ("   11.1 Desglose Modular Funcional", "47"),
        ("   11.2 Interconexión y Flujo de Mensajería Intermodular", "49"),
        ("12. INTERFAZ Y EXPERIENCIA DE USUARIO", "50"),
        ("   12.1 Filosofía de Diseño y Ergonomía Visual", "50"),
        ("   12.2 Experiencia de Escritorio Nativa (CustomTkinter)", "51"),
        ("   12.3 Experiencia Web (Streamlit) y Terminal de Comandos (Rich)", "52"),
        ("13. INFORMACIÓN Y DATOS QUE MANEJA EL SISTEMA", "54"),
        ("   13.1 Variables Fisiológicas y Clínicas Administradas", "54"),
        ("   13.2 Catálogos Estructurados y Almacenes de Reglas", "55"),
        ("14. TECNOLOGÍAS UTILIZADAS", "57"),
        ("   14.1 Ecosistema Python y Selección de Componentes", "57"),
        ("   14.2 Justificación Estratégica de cada Biblioteca", "58"),
        ("15. FUNCIONAMIENTO INTERNO A NIVEL GENERAL", "60"),
        ("   15.1 Arquitectura Desacoplada en Capas", "60"),
        ("   15.2 El Ciclo Cognitivo de Inferencia (Forward Chaining)", "62"),
        ("16. SEGURIDAD Y CONTROL DE ACCESO", "64"),
        ("   16.1 Privacidad Radical y Residencia Local de Datos", "64"),
        ("   16.2 Seguridad Criptográfica y Validación Robusta", "65"),
        ("17. BENEFICIOS DEL PROYECTO", "66"),
        ("   17.1 Impacto Cuantitativo y Cualitativo en el Usuario", "66"),
        ("18. VENTAJAS DEL SISTEMA", "68"),
        ("   18.1 Determinismo y Ausencia de Alucinaciones Computacionales", "68"),
        ("   18.2 Eficiencia de Costos y Autonomía Sin Conexión", "69"),
        ("19. LIMITACIONES ACTUALES", "70"),
        ("   19.1 Evaluación Honesta de Fronteras Operativas", "70"),
        ("20. POSIBLES MEJORAS Y TRABAJO FUTURO", "72"),
        ("   20.1 Hoja de Ruta Tecnológica (Roadmap)", "72"),
        ("21. VALOR DEL PROYECTO", "74"),
        ("   21.1 Trascendencia Sanitaria, Educativa y Científica", "74"),
        ("22. CONCLUSIONES", "76"),
        ("   22.1 Valoración Final del Sistema FitExpert", "76"),
        ("23. GLOSARIO DE TÉRMINOS", "78"),
        ("   23.1 Definición Conceptual de Conceptos Clave", "78"),
        ("24. ANEXOS", "81"),
        ("   Anexo A: Matriz Completa de las 30 Reglas de la Base de Conocimiento", "81"),
        ("   Anexo B: Matriz de Biomecánica y Sustitución de Ejercicios por Lesión", "86"),
        ("   Anexo C: Catálogo de Alérgenos y Sustitutos Culinarios", "88"),
        ("   Anexo D: Estructura de Navegación de Pantallas", "90"),
        ("   Anexo E: Ficha Técnica de Auditoría de Calidad ISO/IEC 25010", "91"),
    ]

    tbl_toc = doc.add_table(rows=len(toc_items), cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_toc.autofit = False

    for idx, (title, page_num) in enumerate(toc_items):
        row_cells = tbl_toc.rows[idx].cells
        row_cells[0].width = Inches(5.7)
        row_cells[1].width = Inches(0.8)

        is_main = not title.startswith("   ")
        shd_col = "FFFFFF" if idx % 2 == 0 else HEX_BG_LIGHT
        for c in range(2):
            shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), shd_col))
            row_cells[c]._tc.get_or_add_tcPr().append(shd)

        # Title
        p_t = row_cells[0].paragraphs[0]
        p_t.paragraph_format.space_before = Pt(2)
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run(title)
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(9.5 if is_main else 8.5)
        r_t.bold = is_main
        r_t.font.color.rgb = COLOR_NAVY_DARK if is_main else COLOR_TEXT_MAIN

        # Page
        p_p = row_cells[1].paragraphs[0]
        p_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_p.paragraph_format.space_before = Pt(2)
        p_p.paragraph_format.space_after = Pt(2)
        r_p = p_p.add_run(page_num)
        r_p.font.name = "Calibri"
        r_p.font.size = Pt(9.5 if is_main else 8.5)
        r_p.bold = is_main
        r_p.font.color.rgb = COLOR_BLUE_ACCENT if is_main else COLOR_TEXT_MUTED

    tblBorders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="0F172A"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="2" w:space="0" w:color="F1F5F9"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tbl_toc._tbl.tblPr.append(tblBorders)

    doc.add_page_break()
